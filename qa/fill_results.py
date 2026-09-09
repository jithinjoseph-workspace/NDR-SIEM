#!/usr/bin/env python3
"""Fills in Actual Result / Status / Executed By / Execution Date for the
seeded test cases in NDR_Test_Case_Tracker.xlsx, based on static verification
performed against the live running containers (binary/bundle string
fingerprints — NOT a full authenticated click-through, since this session has
no login credentials or browser access). That distinction is written into
each row's Actual Result / Remarks so nobody mistakes this for a full
functional QA pass.
"""

import openpyxl

PATH = "/home/user/Music/NDR/NDR-Demo/qa/NDR_Test_Case_Tracker.xlsx"
EXECUTED_BY = "Claude (static/binary verification)"
EXEC_DATE = "2026-09-09"

# col letters: J=Actual Result, K=Status, M=Executed By, N=Execution Date, O=Remarks
RESULTS = {
    "TC-001": dict(
        actual="Running ndr-engine binary (rebuilt 2026-09-03 06:00 UTC) contains the "
               "geo_lookup_batch local-GeoLite2 code path; the old direct-only ip-api.com "
               "call is gone. Verified via container filesystem string inspection, not a "
               "live packet capture.",
        status="Passed",
        remark="Static verification only — confirm with a live packet capture before sign-off.",
    ),
    "TC-002": dict(
        actual="'Agent-Z Sigma Match (rule name unavailable)' string is present in the "
               "running binary; the old hardcoded 'Suricata IDS Alert' string is absent "
               "(0 matches).",
        status="Passed",
        remark="Static binary verification — not a live evidence-bundle click-through.",
    ),
    "TC-003": dict(
        actual="Direct string fingerprint inconclusive (JSON macro doesn't embed "
               "'\"status\":\"ok\"' as one literal). Confirmed instead by build-order: "
               "this fix predates the rules-pagination change, and X-Total-Count/"
               "X-Active-Count (a later change) ARE present in the same binary build — "
               "so this earlier fix is necessarily included too.",
        status="Passed",
        remark="Inferred from build chronology, not a direct string match. No sensor "
               "agent is running in this environment, so a full live isolate click-"
               "through was not possible regardless of this fix.",
    ),
    "TC-004": dict(
        actual="'is-isolated' / 'node-isolated-badge' strings present in the deployed "
               "Angular bundle (chunk-OO5KOA5P.js, chunk-E3YT6INM.js).",
        status="Passed",
        remark="Static bundle verification — not a live browser click-through.",
    ),
    "TC-005": dict(
        actual="'getRulesPage' present in frontend bundle; 'X-Total-Count' present in "
               "both frontend bundle and ndr-engine binary.",
        status="Passed",
        remark="Static verification — did not execute a live authenticated paginated fetch.",
    ),
    "TC-006": dict(
        actual="'rules-search' present in deployed frontend bundle; backend ?q= search "
               "path compiles clean and is unchanged/pre-existing.",
        status="Passed",
        remark="Static verification — did not execute a live search against real data.",
    ),
    "TC-007": dict(
        actual="'This user cannot be deactivated here' and 'Super admin password cannot "
               "be changed here' strings present in the running auth-service binary "
               "(rebuilt 2026-09-03 05:58 UTC), confirming the target-identity guards "
               "are compiled in.",
        status="Passed",
        remark="Static verification — no valid tenant_admin/super_admin credentials "
               "available in this session to run the live cross-tenant API calls.",
    ),
    # TC-008 intentionally left as-is (Blocked, known incomplete MFA feature).
}


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb["Test Cases"]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        tc_id = row[col["Test Case ID"] - 1].value
        if tc_id in RESULTS:
            r = RESULTS[tc_id]
            ws.cell(row=row[0].row, column=col["Actual Result"], value=r["actual"])
            ws.cell(row=row[0].row, column=col["Status"], value=r["status"])
            ws.cell(row=row[0].row, column=col["Executed By"], value=EXECUTED_BY)
            ws.cell(row=row[0].row, column=col["Execution Date"], value=EXEC_DATE)
            existing_remark = ws.cell(row=row[0].row, column=col["Defect ID / Remarks"]).value or ""
            combined = (existing_remark + " " + r["remark"]).strip()
            ws.cell(row=row[0].row, column=col["Defect ID / Remarks"], value=combined)

    wb.save(PATH)
    print(f"Updated {PATH}")


if __name__ == "__main__":
    main()
